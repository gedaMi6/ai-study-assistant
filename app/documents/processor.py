import pymupdf
from pathlib import Path
from pptx import Presentation


def extract_text_from_pdf(file_path: Path) -> str:
    """Extract text from every page of a PDF."""

    document = pymupdf.open(file_path)

    pages = []

    for page in document:
        text = page.get_text()

        if text.strip():
            pages.append(text)

    document.close()

    return "\n\n".join(pages)


def extract_text_from_pptx(file_path: Path) -> str:
    """Extract text from every slide of a PowerPoint presentation."""

    presentation = Presentation(file_path)

    slides = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):
        slide_parts = [f"Slide {slide_number}"]

        for shape in slide.shapes:

            # Normal text boxes, titles, subtitles, etc.
            if hasattr(shape, "text") and shape.text.strip():
                slide_parts.append(shape.text.strip())

            # Tables
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = []

                    for cell in row.cells:
                        text = cell.text.strip()

                        if text:
                            row_text.append(text)

                    if row_text:
                        slide_parts.append(" | ".join(row_text))

        slide_text = "\n".join(slide_parts)

        if len(slide_text.strip()) > len(f"Slide {slide_number}"):
            slides.append(slide_text)

    return "\n\n".join(slides)
